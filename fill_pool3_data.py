#!/usr/bin/env python3
"""
Script to fill missing data in Pool 3 Return Pickup sheet.

This script:
1. Reads the source file with correct customer information
2. Reads the target file Pool 3 sheet that has missing data
3. Matches records by Order Number
4. Fills in the missing columns: Buyer ID, Name, Phone No., District, 
   Category 1, Category 2, Order Value, Product Name
5. Saves the updated file
"""

import pandas as pd
import os
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

# File paths
base_path = "/home/runner/work/datacleaner/datacleaner/need to fix"
source_file = os.path.join(base_path, "Customer Survery Not Touched but correct informations.xlsx")
target_file = os.path.join(base_path, "Customer Survery Partially Completed.xlsx")

# Read source data
print("Reading source file...")
df_source = pd.read_excel(source_file)
print(f"Source file has {len(df_source)} records")

# Create a dictionary for fast lookup by Order Number
source_dict = {}
for _, row in df_source.iterrows():
    order_no = int(row['CustomerOrderCode'])
    source_dict[order_no] = {
        'Buyer ID': int(row['CustomerId']),
        'Name': row['CustomerName'],
        'Phone No.': int(row['CustomerPhone']),
        'District': row['CSState'],
        'Category 1': row['Vertical'],
        'Category 2': row['cat1'],
        'Order Value': int(row['TotalItemPrice']),
        'Product Name': row['ProductName']
    }

print(f"Created lookup dictionary with {len(source_dict)} orders")

# Read target file Pool 3 sheet
print("\nReading target file Pool 3 sheet...")
df_target = pd.read_excel(target_file, sheet_name='Pool 3 Return Pickup', header=3)
print(f"Target sheet has {len(df_target)} records")

# Track statistics
matched_count = 0
unmatched_count = 0
unmatched_orders = []

# Fill in the missing data
print("\nFilling missing data...")
for idx, row in df_target.iterrows():
    order_no = row['Order No.']
    
    if pd.notna(order_no):
        order_no = int(order_no)
        
        # Look up the order in source data
        if order_no in source_dict:
            matched_count += 1
            source_data = source_dict[order_no]
            
            # Fill in the missing columns
            for col, value in source_data.items():
                df_target.at[idx, col] = value
        else:
            unmatched_count += 1
            unmatched_orders.append(order_no)

print(f"\nMatched orders: {matched_count}")
print(f"Unmatched orders: {unmatched_count}")
if unmatched_orders:
    print(f"Unmatched order numbers: {unmatched_orders}")

# Verify the data was filled
print("\nVerifying data population...")
columns_to_check = ['Buyer ID', 'Name', 'Phone No.', 'District', 'Category 1', 'Category 2', 'Order Value', 'Product Name']
for col in columns_to_check:
    null_count = df_target[col].isnull().sum()
    filled_count = df_target[col].notna().sum()
    print(f"{col}: {filled_count} filled, {null_count} still empty")

# Now save the updated data back to the Excel file
# We need to preserve the existing structure including the header rows
print("\nSaving updated file...")

# Load the workbook
wb = load_workbook(target_file)
ws = wb['Pool 3 Return Pickup']

# The data starts at row 5 (after the 4 header rows: 0, 1, 2, 3)
# Update only the columns we modified
column_indices = {
    'Buyer ID': 5,      # Column E (0-indexed: 4, but Excel uses 1-indexed)
    'Name': 6,          # Column F
    'Phone No.': 7,     # Column G
    'District': 8,      # Column H
    'Category 1': 9,    # Column I
    'Category 2': 10,   # Column J
    'Order Value': 11,  # Column K
    'Product Name': 12  # Column L
}

# Start from row 5 (after headers at rows 1-4)
start_row = 5

for df_idx, row in df_target.iterrows():
    excel_row = start_row + df_idx
    
    for col_name, col_idx in column_indices.items():
        value = row[col_name]
        if pd.notna(value):
            ws.cell(row=excel_row, column=col_idx, value=value)

# Save the workbook
wb.save(target_file)
print(f"\nFile saved successfully: {target_file}")

# Final verification - read the file again and check
print("\nFinal verification...")
df_verify = pd.read_excel(target_file, sheet_name='Pool 3 Return Pickup', header=3)
for col in columns_to_check:
    null_count = df_verify[col].isnull().sum()
    filled_count = df_verify[col].notna().sum()
    print(f"{col}: {filled_count} filled, {null_count} empty")

print("\n✓ Data population completed successfully!")
